# -*- coding: utf-8 -*-
"""Génère la présentation PowerPoint ÉpiWatch (16:9) pour le concours INNOVA."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

OUT_DIR = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(OUT_DIR, "EpiWatch_Presentation.pptx")

# Palette
NAVY = RGBColor(0x0F, 0x17, 0x2A)
NAVY2 = RGBColor(0x1E, 0x29, 0x3B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0xB6, 0xC2, 0xD4)
ACCENT = RGBColor(0x2E, 0x6B, 0xE6)
ACCENT_LIGHT = RGBColor(0x60, 0xA5, 0xFA)
GREEN = RGBColor(0x16, 0xA3, 0x4A)
YELLOW = RGBColor(0xEA, 0xB3, 0x08)
ORANGE = RGBColor(0xF9, 0x73, 0x16)
RED = RGBColor(0xEF, 0x44, 0x44)
DARKRED = RGBColor(0x7F, 0x1D, 0x1D)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def slide():
    return prs.slides.add_slide(BLANK)


def bg(s, color):
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color


def rect(s, x, y, w, h, color, line=None, rounded=False):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, x, y, w, h
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if line:
        shp.line.color.rgb = line
        shp.line.width = Pt(1.5)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def circle(s, x, y, d, color):
    shp = s.shapes.add_shape(MSO_SHAPE.OVAL, x, y, d, d)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    if isinstance(runs, str):
        runs = [runs]
    for i, item in enumerate(runs):
        if isinstance(item, str):
            item = {"t": item}
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run()
        r.text = item["t"]
        r.font.size = Pt(item.get("size", 18))
        r.font.bold = item.get("bold", False)
        r.font.color.rgb = item.get("color", WHITE)
        r.font.name = "Segoe UI"
    return tb


def footer(s, n):
    text(s, Inches(0.6), Inches(7.02), Inches(9), Inches(0.35),
         [{"t": "ÉpiWatch  ·  Surveillance épidémiologique en temps réel  ·  INNOVA",
           "size": 10, "color": MUTED}])
    text(s, Inches(12.1), Inches(7.02), Inches(0.7), Inches(0.35),
         [{"t": str(n), "size": 10, "color": MUTED}], align=PP_ALIGN.RIGHT)


def header(s, kicker, title):
    rect(s, 0, 0, SW, Inches(0.14), ACCENT)
    text(s, Inches(0.7), Inches(0.55), Inches(11), Inches(0.4),
         [{"t": kicker.upper(), "size": 13, "bold": True, "color": ACCENT_LIGHT}])
    text(s, Inches(0.7), Inches(0.92), Inches(12), Inches(0.9),
         [{"t": title, "size": 32, "bold": True, "color": WHITE}])


N = 0
def page(kicker, title, fn):
    global N
    N += 1
    s = slide()
    bg(s, NAVY)
    if kicker is not None:
        header(s, kicker, title)
    fn(s)
    footer(s, N)
    return s


# ───────────────────────────── 1. Couverture ─────────────────────────────
def cover(s):
    bg(s, NAVY)
    rect(s, 0, 0, SW, SH, NAVY)
    # bandeau latéral
    rect(s, 0, Inches(2.4), Inches(0.22), Inches(2.7), ACCENT)
    # logo : cercle + croix
    circle(s, Inches(1.0), Inches(2.35), Inches(0.95), ACCENT)
    rect(s, Inches(1.35), Inches(2.55), Inches(0.26), Inches(0.55), WHITE)
    rect(s, Inches(1.21), Inches(2.69), Inches(0.55), Inches(0.26), WHITE)
    text(s, Inches(2.15), Inches(2.42), Inches(10), Inches(0.9),
         [{"t": "ÉpiWatch", "size": 54, "bold": True, "color": WHITE}])
    text(s, Inches(2.2), Inches(3.35), Inches(10.5), Inches(1.0),
         [{"t": "Plateforme de surveillance sanitaire et d'alerte épidémiologique en temps réel",
           "size": 20, "color": MUTED}])
    rect(s, Inches(2.2), Inches(4.35), Inches(3.6), Inches(0.05), ACCENT)
    text(s, Inches(2.2), Inches(4.6), Inches(10.5), Inches(1.2),
         [{"t": "Concours INNOVA", "size": 16, "bold": True, "color": ACCENT_LIGHT},
          {"t": "Prototype fonctionnel · 30 hôpitaux · 5 maladies · 7 niveaux d'alerte",
           "size": 14, "color": MUTED}], spacing=1.3)
    circle(s, Inches(11.6), Inches(1.1), Inches(0.55), GREEN)
    circle(s, Inches(12.15), Inches(1.35), Inches(0.4), YELLOW)
    circle(s, Inches(11.95), Inches(1.75), Inches(0.45), ORANGE)
    circle(s, Inches(12.3), Inches(2.05), Inches(0.5), RED)
    circle(s, Inches(11.85), Inches(2.35), Inches(0.42), DARKRED)


# ───────────────────────────── 2. Problématique ─────────────────────────────
def problem(s):
    text(s, Inches(0.7), Inches(2.0), Inches(11.9), Inches(0.6),
         [{"t": "Les données de santé circulent lentement, restent cloisonnées par "
                "établissement, et les populations ne sont pas alertées à temps.",
           "size": 22, "color": WHITE}], spacing=1.3)
    y = Inches(3.6)
    cards = [
        ("Information cloisonnée", "Chaque hôpital garde ses données pour lui.", RED),
        ("Détection tardive", "Une flambée n'est visible que trop tard.", ORANGE),
        ("Population non informée", "Aucune alerte locale fiable et rapide.", YELLOW),
    ]
    w = Inches(3.7)
    for i, (t, d, c) in enumerate(cards):
        x = Inches(0.7) + i * (w + Inches(0.4))
        rect(s, x, y, w, Inches(2.2), NAVY2, rounded=True)
        rect(s, x, y, w, Inches(0.09), c)
        text(s, x + Inches(0.35), y + Inches(0.45), w - Inches(0.7), Inches(1.6),
             [{"t": t, "size": 19, "bold": True, "color": WHITE},
              {"t": d, "size": 15, "color": MUTED}], spacing=1.2)


# ───────────────────────────── 3. Solution ─────────────────────────────
def solution(s):
    text(s, Inches(0.7), Inches(2.0), Inches(11.9), Inches(0.7),
         [{"t": "ÉpiWatch centralise, détecte, visualise et alerte — en temps réel.",
           "size": 22, "bold": True, "color": ACCENT_LIGHT}])
    steps = [
        ("1", "Saisie", "L'agent enregistre les cas depuis l'hôpital."),
        ("2", "Centralisation", "Toutes les données dans une base commune."),
        ("3", "Analyse & carte", "Le niveau d'alerte s'affiche sur la carte."),
        ("4", "Alerte", "Responsables et population sont prévenus."),
    ]
    y = Inches(3.2)
    w = Inches(2.75)
    for i, (n, t, d) in enumerate(steps):
        x = Inches(0.7) + i * (w + Inches(0.35))
        rect(s, x, y, w, Inches(2.7), NAVY2, rounded=True)
        circle(s, x + Inches(0.3), y + Inches(0.3), Inches(0.65), ACCENT)
        text(s, x + Inches(0.3), y + Inches(0.33), Inches(0.65), Inches(0.6),
             [{"t": n, "size": 22, "bold": True, "color": WHITE}], align=PP_ALIGN.CENTER)
        text(s, x + Inches(0.3), y + Inches(1.15), w - Inches(0.6), Inches(1.4),
             [{"t": t, "size": 18, "bold": True, "color": WHITE},
              {"t": d, "size": 13, "color": MUTED}], spacing=1.15)
        if i < 3:
            text(s, x + w + Inches(0.02), y + Inches(1.0), Inches(0.35), Inches(0.6),
                 [{"t": "›", "size": 30, "bold": True, "color": ACCENT}], align=PP_ALIGN.CENTER)


# ───────────────────────────── 4. Niveaux d'alerte ─────────────────────────────
def levels(s):
    data = [
        (0, "🔵", "Bleu", "Surveillance habituelle", RGBColor(0x25, 0x63, 0xEB)),
        (1, "🟢", "Vert", "Situation stable", GREEN),
        (2, "🟡", "Jaune", "Signal à surveiller", YELLOW),
        (3, "🟠", "Orange", "Hausse préoccupante", ORANGE),
        (4, "🔴", "Rouge léger", "Alerte élevée", RED),
        (5, "🔴", "Rouge", "Situation critique", RGBColor(0xB9, 0x1C, 0x1C)),
        (6, "🟥", "Rouge foncé", "Alerte majeure", DARKRED),
    ]
    y = Inches(2.05)
    h = Inches(0.62)
    for i, (n, emo, name, desc, c) in enumerate(data):
        yy = y + i * (h + Inches(0.08))
        rect(s, Inches(0.7), yy, Inches(11.9), h, NAVY2, rounded=True)
        rect(s, Inches(0.7), yy, Inches(0.14), h, c)
        circle(s, Inches(1.1), yy + Inches(0.09), Inches(0.44), c)
        text(s, Inches(1.1), yy + Inches(0.1), Inches(0.44), Inches(0.42),
             [{"t": str(n), "size": 16, "bold": True, "color": WHITE}], align=PP_ALIGN.CENTER)
        text(s, Inches(1.8), yy + Inches(0.13), Inches(3), Inches(0.4),
             [{"t": name, "size": 16, "bold": True, "color": WHITE}])
        text(s, Inches(5.0), yy + Inches(0.15), Inches(7.4), Inches(0.4),
             [{"t": desc, "size": 15, "color": MUTED}])


# ───────────────────────────── 5. Escalade ─────────────────────────────
def escalation(s):
    y = Inches(2.5)
    boxes = [
        ("HÔPITAL", "Un hôpital franchit un seuil", "Alerte locale", ACCENT),
        ("DISTRICT", "≥ 2 hôpitaux au niveau 5+", "Alerte de district", ORANGE),
        ("RÉGION", "≥ 5 hôpitaux au niveau 5+", "Alerte régionale", DARKRED),
    ]
    w = Inches(3.5)
    for i, (t, d, a, c) in enumerate(boxes):
        x = Inches(0.9) + i * (w + Inches(0.7))
        rect(s, x, y, w, Inches(2.9), NAVY2, rounded=True)
        rect(s, x, y, w, Inches(0.12), c)
        text(s, x + Inches(0.35), y + Inches(0.5), w - Inches(0.7), Inches(0.5),
             [{"t": t, "size": 22, "bold": True, "color": c}])
        text(s, x + Inches(0.35), y + Inches(1.15), w - Inches(0.7), Inches(1.0),
             [{"t": d, "size": 15, "color": MUTED}], spacing=1.2)
        text(s, x + Inches(0.35), y + Inches(2.15), w - Inches(0.7), Inches(0.5),
             [{"t": a, "size": 15, "bold": True, "color": WHITE}])
        if i < 2:
            text(s, x + w + Inches(0.05), y + Inches(1.1), Inches(0.6), Inches(0.6),
                 [{"t": "→", "size": 34, "bold": True, "color": ACCENT}], align=PP_ALIGN.CENTER)
    text(s, Inches(0.9), Inches(5.85), Inches(11.5), Inches(0.6),
         [{"t": "Règles entièrement configurables · comptage des établissements distincts · "
                "escalade immédiate au niveau 6", "size": 14, "color": ACCENT_LIGHT}],
         align=PP_ALIGN.CENTER)


# ───────────────────────────── 6. Fonctionnalités ─────────────────────────────
def features(s):
    groups = [
        ("Tableau de bord", ["Indicateurs clés", "Graphiques d'évolution", "Rapports & export CSV"], ACCENT),
        ("Carte interactive", ["30 hôpitaux", "Marqueurs par niveau", "Filtres & légende"], GREEN),
        ("Saisie hospitalière", ["Enregistrement rapide", "Correction tracée", "Historique"], ORANGE),
        ("Alertes & notifications", ["Temps réel (SSE)", "Zones suivies", "Conseils de prévention"], RED),
    ]
    y = Inches(2.1)
    w = Inches(2.9)
    for i, (t, items, c) in enumerate(groups):
        x = Inches(0.7) + i * (w + Inches(0.33))
        rect(s, x, y, w, Inches(4.3), NAVY2, rounded=True)
        rect(s, x, y, w, Inches(0.9), c)
        text(s, x + Inches(0.25), y + Inches(0.2), w - Inches(0.5), Inches(0.6),
             [{"t": t, "size": 16, "bold": True, "color": WHITE}], anchor=MSO_ANCHOR.MIDDLE)
        runs = []
        for it in items:
            runs.append({"t": "•  " + it, "size": 14, "color": MUTED})
        text(s, x + Inches(0.3), y + Inches(1.2), w - Inches(0.6), Inches(2.9),
             runs, spacing=1.9)


# ───────────────────────────── 7. Architecture ─────────────────────────────
def architecture(s):
    y = Inches(2.15)
    # Clients
    rect(s, Inches(0.7), y, Inches(11.9), Inches(1.15), NAVY2, rounded=True)
    text(s, Inches(1.0), y + Inches(0.18), Inches(11), Inches(0.4),
         [{"t": "CLIENTS  —  PWA React", "size": 14, "bold": True, "color": ACCENT_LIGHT}])
    text(s, Inches(1.0), y + Inches(0.6), Inches(11), Inches(0.4),
         [{"t": "Desktop (saisie)  ·  Mobile (alertes)  ·  Tablette     |     "
                "React 18 · Vite · Leaflet · Recharts", "size": 13, "color": MUTED}])
    # flèches
    text(s, Inches(3.0), y + Inches(1.2), Inches(2), Inches(0.5),
         [{"t": "REST + SSE ▼", "size": 13, "bold": True, "color": ACCENT}])
    # Serveur
    ys = y + Inches(1.75)
    rect(s, Inches(0.7), ys, Inches(11.9), Inches(1.15), NAVY2, rounded=True)
    text(s, Inches(1.0), ys + Inches(0.18), Inches(11), Inches(0.4),
         [{"t": "SERVEUR  —  Node.js + Express", "size": 14, "bold": True, "color": ACCENT_LIGHT}])
    text(s, Inches(1.0), ys + Inches(0.6), Inches(11), Inches(0.4),
         [{"t": "Authentification · API REST · Moteur d'alertes · Bus temps réel (SSE) · "
                "Service des fichiers", "size": 13, "color": MUTED}])
    text(s, Inches(3.0), ys + Inches(1.2), Inches(2), Inches(0.5),
         [{"t": "node:sqlite ▼", "size": 13, "bold": True, "color": ACCENT}])
    # Base
    yb = ys + Inches(1.75)
    rect(s, Inches(0.7), yb, Inches(11.9), Inches(0.9), NAVY2, rounded=True)
    text(s, Inches(1.0), yb + Inches(0.12), Inches(11), Inches(0.4),
         [{"t": "BASE  —  SQLite embarquée (WAL)", "size": 14, "bold": True, "color": ACCENT_LIGHT}])
    text(s, Inches(1.0), yb + Inches(0.48), Inches(11), Inches(0.4),
         [{"t": "10 tables · fichier unique · aucune dépendance native · portable & hors-ligne",
           "size": 13, "color": MUTED}])


# ───────────────────────────── 8. Sécurité ─────────────────────────────
def security(s):
    items = [
        ("Rôles (RBAC)", "5 rôles : agent, responsable, autorité, admin, public."),
        ("Cloisonnement", "Un agent ne saisit que pour son hôpital."),
        ("Mots de passe", "Hachage scrypt + jetons signés HMAC."),
        ("Traçabilité", "Journal d'audit de chaque action sensible."),
        ("Confidentialité", "Comptages agrégés, aucune donnée personnelle."),
        ("Hors-ligne", "Fonctionne sur réseau local, sans Internet."),
    ]
    y = Inches(2.1)
    w = Inches(5.85)
    h = Inches(1.35)
    for i, (t, d) in enumerate(items):
        col = i % 2
        row = i // 2
        x = Inches(0.7) + col * (w + Inches(0.3))
        yy = y + row * (h + Inches(0.2))
        rect(s, x, yy, w, h, NAVY2, rounded=True)
        rect(s, x, yy, Inches(0.12), h, ACCENT)
        text(s, x + Inches(0.4), yy + Inches(0.2), w - Inches(0.7), Inches(0.5),
             [{"t": t, "size": 17, "bold": True, "color": WHITE}])
        text(s, x + Inches(0.4), yy + Inches(0.72), w - Inches(0.7), Inches(0.5),
             [{"t": d, "size": 13, "color": MUTED}])


# ───────────────────────────── 9. Démo / chiffres ─────────────────────────────
def demo(s):
    stats = [
        ("30", "hôpitaux réels", GREEN),
        ("5", "maladies suivies", ACCENT),
        ("7", "niveaux d'alerte", ORANGE),
        ("8 400", "déclarations de démo", YELLOW),
        ("38", "tests validés", RED),
    ]
    y = Inches(2.2)
    w = Inches(2.35)
    for i, (n, d, c) in enumerate(stats):
        x = Inches(0.7) + i * (w + Inches(0.17))
        rect(s, x, y, w, Inches(1.9), NAVY2, rounded=True)
        rect(s, x, y, w, Inches(0.1), c)
        text(s, x + Inches(0.15), y + Inches(0.45), w - Inches(0.3), Inches(0.8),
             [{"t": n, "size": 34, "bold": True, "color": c}], align=PP_ALIGN.CENTER)
        text(s, x + Inches(0.15), y + Inches(1.25), w - Inches(0.3), Inches(0.5),
             [{"t": d, "size": 13, "color": MUTED}], align=PP_ALIGN.CENTER)
    text(s, Inches(0.7), Inches(4.6), Inches(11.9), Inches(1.6),
         [{"t": "Scénario de démonstration", "size": 18, "bold": True, "color": ACCENT_LIGHT},
          {"t": "Saisie d'un cas  →  hausse du marqueur  →  alerte locale  →  escalade district "
                "puis région  →  notification sur le téléphone citoyen.", "size": 15, "color": MUTED}],
         spacing=1.3)


# ───────────────────────────── 10. Conclusion ─────────────────────────────
def conclusion(s):
    text(s, Inches(1.2), Inches(2.3), Inches(11), Inches(1.0),
         [{"t": "Un prototype qui fonctionne réellement,", "size": 32, "bold": True, "color": WHITE},
          {"t": "pas une simple maquette.", "size": 32, "bold": True, "color": ACCENT_LIGHT}],
         align=PP_ALIGN.CENTER, spacing=1.1)
    rect(s, Inches(5.6), Inches(4.3), Inches(2.1), Inches(0.05), ACCENT)
    text(s, Inches(1.2), Inches(4.7), Inches(11), Inches(1.6),
         [{"t": "Merci de votre attention", "size": 22, "color": MUTED},
          {"t": "ÉpiWatch  ·  Concours INNOVA", "size": 16, "color": ACCENT_LIGHT}],
         align=PP_ALIGN.CENTER, spacing=1.5)


# Construction
cover(slide())
bg(prs.slides[0], NAVY)
N = 1  # couverture comptée
footer(prs.slides[0], 1)

page("Problématique", "Pourquoi ÉpiWatch ?", problem)
page("Solution", "Ce que fait ÉpiWatch", solution)
page("Niveaux d'alerte", "Sept niveaux, une action pour chacun", levels)
page("Escalade", "De l'hôpital à la région", escalation)
page("Fonctionnalités", "Quatre interfaces, une PWA", features)
page("Architecture", "Une pile simple et robuste", architecture)
page("Sécurité", "Conçu pour la confiance", security)
page("Démonstration", "ÉpiWatch en chiffres", demo)
page(None, None, conclusion)

prs.save(OUT)
print("Présentation générée :", OUT)
print("Nombre de diapositives :", len(prs.slides._sldIdLst))
